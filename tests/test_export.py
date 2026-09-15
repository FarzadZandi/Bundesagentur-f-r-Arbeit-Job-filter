from datetime import date

from openpyxl import load_workbook

from jobfilter.adapters.base import JobAd
from jobfilter.export import (
    DROPPED_COLUMNS,
    KEPT_COLUMNS,
    ScoredJob,
    consolidate_existing_outputs,
    write_outputs,
)
from jobfilter.scoring import ScoreResult


def item(
    band, score, title, *, reason=None, stage="",
    primary="B06 Data Analytics und Business Intelligence",
    employment_format="WERKSTUDENT / WORKING STUDENT",
    location="Heilbronn",
    posted="2026-08-25",
    url=None,
):
    ad = JobAd(
        url=url or f"https://example.com/{title.lower().replace(' ', '-')}",
        title=title,
        employer="Example GmbH",
        location=location,
        posted_date=posted,
        deadline="2026-09-30",
        full_text="KPI reporting and data analysis using Power BI.",
        employment_type="Werkstudent",
    )
    result = ScoreResult(
        score=score,
        raw_score=score,
        band=band,
        reason=reason,
        drop_stage=stage,
        matched_term="python" if band == "DROP" else "",
        employment_format=employment_format,
        primary_cv_family=primary,
        matched_terms={
            "B06.strong": ("data analytics",),
            "analysis_task": ("kpi reporting",),
            "methods": (),
            "tools_strong": ("power bi",),
        },
        language_flags=("GERMAN C1+ / NATIVE-LEVEL",) if title == "Review Job" else (),
        eligibility_flags=("IMMATRIKULATION REQUIRED",),
        penalties=("GERMAN C1+ / NATIVE-LEVEL (-5)",) if title == "Review Job" else (),
        bonuses=("EXACT STUDENT FORMAT (+2)",),
    )
    return ScoredJob(ad, "arbeitsagentur", result)


def test_writes_only_required_workbooks(tmp_path):
    jobs = [
        item("REVIEW", 8, "Review Job"),
        item("PRIORITY", 18, "Priority Job"),
        item("DE REQUIRED", 16, "German Required Job", reason="german_c1_c2_or_native_required"),
        item("ACADEMISCHE", 0, "Academic Job", reason="excluded_doctoral_candidate_role"),
        item("DROP", 0, "Dropped Job", reason="no_target_family_or_analysis_match", stage="relevance_gate", primary=""),
    ]
    full_time_xlsx, other_xlsx = write_outputs(jobs, tmp_path, date(2026, 8, 26))
    assert full_time_xlsx.name == "jobs_vollzeit_2026-08-26.xlsx"
    assert other_xlsx.name == "jobs_other_2026-08-26.xlsx"
    assert full_time_xlsx.parent.name == "run"
    assert other_xlsx.parent.name == "run"
    assert (tmp_path / "jobs_vollzeit.xlsx").exists()
    assert (tmp_path / "jobs_other.xlsx").exists()
    workbook = load_workbook(other_xlsx)
    assert workbook.sheetnames == ["PRIORITY", "REVIEW", "DE Required", "Academische", "DROPPED"]
    assert [cell.value for cell in workbook["PRIORITY"][1]] == KEPT_COLUMNS
    assert [cell.value for cell in workbook["DROPPED"][1]] == DROPPED_COLUMNS
    assert workbook["PRIORITY"].freeze_panes == "A2"
    assert workbook["PRIORITY"].auto_filter.ref == "A1:AC2"
    assert workbook["PRIORITY"]["N2"].hyperlink.target.endswith("priority-job")
    assert workbook["PRIORITY"]["I2"].value == "Baden-Württemberg"
    assert workbook["PRIORITY"]["J2"].value == 0
    assert workbook["REVIEW"]["S2"].value == "GERMAN C1+ / NATIVE-LEVEL"
    assert workbook["DE Required"]["F2"].value == "German Required Job"
    assert workbook["Academische"]["F2"].value == "Academic Job"
    assert workbook["DROPPED"]["H2"].value == "relevance_gate"
    assert workbook["DROPPED"]["I2"].value == "no_target_family_or_analysis_match"
    assert workbook["DROPPED"]["J2"].value == "python"
    assert workbook["PRIORITY"]["Z2"].value == "2.3-2026-09-11"
    assert workbook["PRIORITY"]["AA2"].value == "2026-08-26"
    assert workbook["PRIORITY"]["AB2"].value == "Werkstudent"
    assert "KPI reporting" in workbook["PRIORITY"]["AC2"].value
    assert all(workbook["PRIORITY"].column_dimensions[column].hidden for column in ("Z", "AA", "AB", "AC"))
    assert not list(tmp_path.rglob("*.txt"))


def test_rows_are_sorted_descending(tmp_path):
    _, other_xlsx = write_outputs(
        [item("PRIORITY", 12, "Lower"), item("PRIORITY", 22, "Higher")],
        tmp_path,
        date(2026, 8, 26),
    )
    sheet = load_workbook(other_xlsx)["PRIORITY"]
    assert [sheet["A2"].value, sheet["A3"].value] == [22, 12]


def test_rows_are_sorted_by_posting_date_before_score(tmp_path):
    _, other_xlsx = write_outputs(
        [
            item("PRIORITY", 30, "Older", posted="2026-08-20"),
            item("PRIORITY", 12, "Newer", posted="25.08.2026"),
        ],
        tmp_path,
        date(2026, 8, 26),
    )
    sheet = load_workbook(other_xlsx)["PRIORITY"]
    assert [sheet["F2"].value, sheet["F3"].value] == ["Newer", "Older"]


def test_full_time_and_other_jobs_are_written_to_separate_workbooks(tmp_path):
    full_time_xlsx, other_xlsx = write_outputs(
        [
            item("PRIORITY", 20, "Full-time Analyst", employment_format="FULL-TIME / VOLLZEIT"),
            item("PRIORITY", 18, "Working Student Analyst"),
        ],
        tmp_path,
        date(2026, 8, 26),
    )
    full_time = load_workbook(full_time_xlsx)
    other = load_workbook(other_xlsx)
    assert full_time["PRIORITY"]["F2"].value == "Full-time Analyst"
    assert full_time["PRIORITY"].max_row == 2
    assert other["PRIORITY"]["F2"].value == "Working Student Analyst"
    assert other["PRIORITY"].max_row == 2


def test_repeated_same_day_runs_do_not_overwrite_prior_outputs(tmp_path):
    first = write_outputs([item("PRIORITY", 18, "First Batch")], tmp_path, date(2026, 8, 26))
    second = write_outputs([item("PRIORITY", 18, "Second Batch")], tmp_path, date(2026, 8, 26))
    assert first[0].name == "jobs_vollzeit_2026-08-26.xlsx"
    assert second[0].name == "jobs_vollzeit_2026-08-26_2.xlsx"
    assert all(path.exists() for path in (*first, *second))


def test_master_workbook_deduplicates_urls_and_keeps_latest_record(tmp_path):
    write_outputs(
        [item("REVIEW", 8, "Same Job", posted="2026-08-20")],
        tmp_path,
        date(2026, 8, 26),
    )
    write_outputs(
        [item("PRIORITY", 18, "Same Job", posted="2026-08-25")],
        tmp_path,
        date(2026, 8, 27),
    )
    master = load_workbook(tmp_path / "jobs_other.xlsx")
    assert master["REVIEW"].max_row == 1
    assert master["PRIORITY"].max_row == 2
    assert master["PRIORITY"]["F2"].value == "Same Job"
    assert master["PRIORITY"]["K2"].value == "2026-08-25"


def test_consolidation_moves_legacy_daily_workbooks_into_run_folder(tmp_path):
    full_time, other = write_outputs(
        [
            item("PRIORITY", 20, "Full Time", employment_format="FULL-TIME / VOLLZEIT"),
            item("PRIORITY", 18, "Student"),
        ],
        tmp_path,
        date(2026, 8, 26),
    )
    full_time.replace(tmp_path / full_time.name)
    other.replace(tmp_path / other.name)
    (tmp_path / "jobs_vollzeit.xlsx").unlink()
    (tmp_path / "jobs_other.xlsx").unlink()

    masters = consolidate_existing_outputs(tmp_path)

    assert all(path.exists() for path in masters)
    assert (tmp_path / "run" / full_time.name).exists()
    assert (tmp_path / "run" / other.name).exists()
    assert not (tmp_path / full_time.name).exists()
    assert not (tmp_path / other.name).exists()


def test_cross_master_duplicate_is_kept_only_in_latest_canonical_group(tmp_path):
    shared_url = "https://example.com/jobs/one?utm_source=old"
    write_outputs(
        [item("PRIORITY", 18, "Analyst", employment_format="FULL-TIME / VOLLZEIT", posted="2026-08-20", url=shared_url)],
        tmp_path,
        date(2026, 8, 26),
    )
    write_outputs(
        [item("PRIORITY", 20, "Working Student Analyst", posted="2026-08-25", url="https://example.com/jobs/one")],
        tmp_path,
        date(2026, 8, 27),
    )
    full_time = load_workbook(tmp_path / "jobs_vollzeit.xlsx")
    other = load_workbook(tmp_path / "jobs_other.xlsx")
    assert full_time["PRIORITY"].max_row == 1
    assert other["PRIORITY"].max_row == 2
    assert other["PRIORITY"]["F2"].value == "Working Student Analyst"


def test_duplicates_are_removed_before_daily_split(tmp_path):
    shared_url = "https://example.com/jobs/two"
    full_path, other_path = write_outputs(
        [
            item("PRIORITY", 18, "Old Full Time", employment_format="FULL-TIME / VOLLZEIT", posted="2026-08-20", url=shared_url),
            item("PRIORITY", 20, "New Student", posted="2026-08-25", url=shared_url),
        ],
        tmp_path,
        date(2026, 8, 26),
    )
    assert load_workbook(full_path)["PRIORITY"].max_row == 1
    assert load_workbook(other_path)["PRIORITY"]["F2"].value == "New Student"
