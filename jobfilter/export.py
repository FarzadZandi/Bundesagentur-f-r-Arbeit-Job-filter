from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import re
import shutil
from typing import Any, Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .adapters.base import JobAd
from .geography import enrich_location
from .scoring import (
    SCORING_RULE_VERSION,
    ScoreResult,
    load_keywords,
    priority_title_hits,
    score_text,
    title_gate,
)
from .store import canonicalize_url


KEPT_COLUMNS = [
    "score", "band", "employment_format", "primary_cv_family", "secondary_cv_family",
    "title", "employer", "location", "federal_state", "distance_from_heilbronn_km",
    "posted", "deadline", "board", "url",
    "family_hits", "analysis_task_hits", "method_hits", "tool_hits", "language_flags",
    "eligibility_flags", "other_flags", "penalties", "bonuses", "kill_reason", "snippet",
    "rule_version", "fetched_date", "raw_employment_type", "full_text",
]
DROPPED_COLUMNS = [
    "title", "employer", "location", "federal_state", "distance_from_heilbronn_km",
    "posted", "url", "drop_stage", "drop_reason", "matched_term",
    "score", "employment_format", "board", "rule_version", "fetched_date",
    "raw_employment_type", "full_text", "snippet", "family_hits", "analysis_task_hits",
    "language_flags",
]
WIDTHS = [8, 11, 25, 40, 40, 48, 32, 23, 24, 27, 13, 20, 18, 48, 45, 40, 35, 42, 32, 36, 40, 32, 28, 32, 60, 22, 14, 24, 90]
DROPPED_WIDTHS = [48, 32, 23, 24, 27, 13, 48, 22, 40, 30, 8, 25, 18, 22, 14, 24, 90, 60, 45, 40, 32]
HEADER_FILL = PatternFill("solid", fgColor="174A6E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PRIORITY_FILL = PatternFill("solid", fgColor="D9EAF7")
AMBER_FILL = PatternFill("solid", fgColor="FCE4A6")
SHEET_NAMES = ("PRIORITY", "REVIEW", "DE Required", "Academische", "DROPPED")
DAILY_WORKBOOK_PATTERN = re.compile(
    r"^jobs_(vollzeit|other)_(\d{4}-\d{2}-\d{2})(?:_(\d+))?\.xlsx$",
    re.IGNORECASE,
)
DEFAULT_KEYWORDS_PATH = Path(__file__).resolve().parent.parent / "keywords_jobseeking_student_v1.json"


@dataclass(frozen=True)
class ScoredJob:
    ad: JobAd
    board: str
    result: ScoreResult


def _join(values: Iterable[str]) -> str:
    return "; ".join(dict.fromkeys(value for value in values if value))


def _hits(result: ScoreResult, predicate) -> str:
    values: list[str] = []
    for group, hits in result.matched_terms.items():
        if predicate(group):
            values.extend(hits)
    return _join(values)


def _kept_row(item: ScoredJob, fetched_date: date) -> list[object]:
    ad, result = item.ad, item.result
    family_hits = _hits(result, lambda group: group.startswith("B") and group.endswith((".strong", ".weak")))
    location_info = enrich_location(ad.location)
    return [
        result.score, result.band, result.employment_format, result.primary_cv_family,
        result.secondary_cv_family, ad.title, ad.employer, ad.location,
        location_info.federal_state, location_info.distance_from_heilbronn_km,
        ad.posted_date, ad.deadline, item.board, ad.url, family_hits,
        _join(result.matched_terms.get("analysis_task", ())),
        _join(result.matched_terms.get("methods", ())),
        _join((*result.matched_terms.get("tools_strong", ()), *result.matched_terms.get("tools_basic_or_adjacent", ()))),
        _join(result.language_flags), _join(result.eligibility_flags), _join(result.other_flags),
        _join(result.penalties), _join(result.bonuses), result.reason or "",
        " ".join(ad.full_text.split())[:500],
        SCORING_RULE_VERSION, fetched_date.isoformat(), ad.employment_type,
        " ".join(ad.full_text.split())[:32000],
    ]


def _style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 28
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 3)
    headers = {cell.value: cell.column for cell in sheet[1]}
    if "distance_from_heilbronn_km" in headers:
        column = get_column_letter(headers["distance_from_heilbronn_km"])
        for cell in sheet[column][1:]:
            cell.number_format = "0.0"
    for internal_header in ("rule_version", "fetched_date", "raw_employment_type", "full_text"):
        if internal_header in headers:
            sheet.column_dimensions[get_column_letter(headers[internal_header])].hidden = True


def _parse_posted(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def _item_sort_key(item: ScoredJob) -> tuple[date, int]:
    return _parse_posted(item.ad.posted_date) or date.min, item.result.score


def _write_workbook(items: list[ScoredJob], xlsx_path: Path, fetched_date: date) -> None:
    by_band = {
        band: sorted(
            (item for item in items if item.result.band == band),
            key=_item_sort_key,
            reverse=True,
        )
        for band in ("PRIORITY", "REVIEW", "DE REQUIRED", "ACADEMISCHE", "DROP")
    }
    workbook = Workbook()
    workbook.remove(workbook.active)
    for band, sheet_name in (
        ("PRIORITY", "PRIORITY"),
        ("REVIEW", "REVIEW"),
        ("DE REQUIRED", "DE Required"),
        ("ACADEMISCHE", "Academische"),
    ):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(KEPT_COLUMNS)
        for item in by_band[band]:
            sheet.append(_kept_row(item, fetched_date))
            row = sheet.max_row
            url_column = KEPT_COLUMNS.index("url") + 1
            sheet.cell(row, url_column).hyperlink = item.ad.url
            sheet.cell(row, url_column).style = "Hyperlink"
            if band == "PRIORITY":
                sheet.cell(row, 2).fill = PRIORITY_FILL
            if band == "DE REQUIRED" or any("GERMAN C1" in value for value in item.result.language_flags):
                sheet.cell(row, KEPT_COLUMNS.index("language_flags") + 1).fill = AMBER_FILL
        _style_sheet(sheet, WIDTHS)

    dropped = workbook.create_sheet("DROPPED")
    dropped.append(DROPPED_COLUMNS)
    for item in by_band["DROP"]:
        result = item.result
        location_info = enrich_location(item.ad.location)
        dropped.append([
            item.ad.title, item.ad.employer, item.ad.location,
            location_info.federal_state, location_info.distance_from_heilbronn_km,
            item.ad.posted_date, item.ad.url, result.drop_stage,
            result.reason or "score_below_review", result.matched_term,
            result.score, result.employment_format, item.board, SCORING_RULE_VERSION,
            fetched_date.isoformat(), item.ad.employment_type,
            " ".join(item.ad.full_text.split())[:32000],
            " ".join(item.ad.full_text.split())[:500],
            _hits(result, lambda group: group.startswith("B") and group.endswith((".strong", ".weak"))),
            _join(result.matched_terms.get("analysis_task", ())),
            _join(result.language_flags),
        ])
        url_column = DROPPED_COLUMNS.index("url") + 1
        dropped.cell(dropped.max_row, url_column).hyperlink = item.ad.url
        dropped.cell(dropped.max_row, url_column).style = "Hyperlink"
    _style_sheet(dropped, DROPPED_WIDTHS)

    workbook.save(xlsx_path)


def _record_key(record: Mapping[str, object]) -> str:
    url = str(record.get("url") or "").strip()
    if url:
        return f"url:{canonicalize_url(url)}"
    identity = "\n".join(
        " ".join(str(record.get(column) or "").casefold().split())
        for column in ("title", "employer", "location")
    )
    return f"identity:{identity}"


def _read_records(path: Path, master_group: str = "") -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    records: list[dict[str, object]] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.title not in SHEET_NAMES:
                continue
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if not any(value not in (None, "") for value in values):
                    continue
                record = {
                    str(header): value
                    for header, value in zip(headers, values)
                    if header not in (None, "")
                }
                location_info = enrich_location(str(record.get("location") or ""))
                record["federal_state"] = location_info.federal_state
                record["distance_from_heilbronn_km"] = location_info.distance_from_heilbronn_km
                record["_sheet"] = sheet.title
                record["_master_group"] = master_group
                records.append(record)
    finally:
        workbook.close()
    return records


def _deduplicate_records(records: Iterable[dict[str, object]]) -> list[dict[str, object]]:
    unique: dict[str, dict[str, object]] = {}
    for record in records:
        key = _record_key(record)
        existing = unique.get(key)
        if existing is None:
            unique[key] = dict(record)
            continue
        incoming_wins = _record_freshness_key(record) >= _record_freshness_key(existing)
        winner, loser = (record, existing) if incoming_wins else (existing, record)
        combined = dict(loser)
        for column, value in winner.items():
            if value not in (None, ""):
                combined[column] = value
        combined["_sheet"] = winner["_sheet"]
        combined["_master_group"] = winner.get("_master_group", "")
        unique[key] = combined
    return list(unique.values())


def _record_sort_key(record: Mapping[str, object]) -> tuple[date, float]:
    posted = _parse_posted(record.get("posted")) or date.min
    try:
        score = float(record.get("score") or 0)
    except (TypeError, ValueError):
        score = 0
    return posted, score


def _record_freshness_key(record: Mapping[str, object]) -> tuple[date, date, float]:
    fetched = _parse_posted(record.get("fetched_date")) or date.min
    posted, score = _record_sort_key(record)
    return fetched, posted, score


def _sheet_for_band(band: str) -> str:
    return {
        "PRIORITY": "PRIORITY",
        "REVIEW": "REVIEW",
        "DE REQUIRED": "DE Required",
        "ACADEMISCHE": "Academische",
        "DROP": "DROPPED",
    }[band]


def _apply_result_to_record(record: dict[str, object], result: ScoreResult) -> None:
    record.update({
        "score": result.score,
        "band": result.band,
        "employment_format": result.employment_format or record.get("employment_format", ""),
        "primary_cv_family": result.primary_cv_family,
        "secondary_cv_family": result.secondary_cv_family,
        "family_hits": _hits(result, lambda group: group.startswith("B") and group.endswith((".strong", ".weak"))),
        "analysis_task_hits": _join(result.matched_terms.get("analysis_task", ())),
        "method_hits": _join(result.matched_terms.get("methods", ())),
        "tool_hits": _join((*result.matched_terms.get("tools_strong", ()), *result.matched_terms.get("tools_basic_or_adjacent", ()))),
        "language_flags": _join(result.language_flags),
        "eligibility_flags": _join(result.eligibility_flags),
        "other_flags": _join(result.other_flags),
        "penalties": _join(result.penalties),
        "bonuses": _join(result.bonuses),
        "kill_reason": result.reason or "",
        "drop_stage": result.drop_stage,
        "drop_reason": result.reason or ("score_below_review" if result.band == "DROP" else ""),
        "matched_term": result.matched_term,
        "_sheet": _sheet_for_band(result.band),
    })


def _reclassify_record(record: dict[str, object], keywords: Mapping[str, Any]) -> dict[str, object]:
    """Reapply current rules when enough historical evidence exists.

    Legacy rows without a stored description can still receive deterministic
    title gates and Priority demotion, but are never promoted from body evidence.
    """
    record = dict(record)
    if str(record.get("rule_version") or "") == SCORING_RULE_VERSION:
        return record
    title = str(record.get("title") or "")
    full_text = str(record.get("full_text") or "")
    raw_employment = str(record.get("raw_employment_type") or record.get("employment_format") or "")
    if full_text:
        result = score_text(title, full_text, keywords, employment_type=raw_employment)
        _apply_result_to_record(record, result)
        record["rule_version"] = SCORING_RULE_VERSION
        return record

    deterministic_gate = title_gate(title, keywords)
    if deterministic_gate:
        band, reason, matched_term = deterministic_gate
        record.update({
            "band": band,
            "score": 0,
            "kill_reason": reason,
            "drop_reason": reason if band == "DROP" else "",
            "drop_stage": (
                "seniority_gate" if reason == "excluded_senior_or_leadership_role"
                else "commercial_role_gate" if reason == "excluded_einkauf_seller_or_sales_role"
                else "pure_it_title_gate" if band == "DROP" else ""
            ),
            "matched_term": matched_term,
            "_sheet": _sheet_for_band(band),
        })
    elif str(record.get("_sheet") or "") == "PRIORITY" and not priority_title_hits(title, keywords):
        record.update({
            "band": "REVIEW",
            "kill_reason": "priority_requires_target_title_evidence",
            "_sheet": "REVIEW",
        })
    record["rule_version"] = f"{SCORING_RULE_VERSION}-title-only-legacy"
    return record


def _write_records_workbook(records: Iterable[dict[str, object]], xlsx_path: Path) -> None:
    grouped = {sheet_name: [] for sheet_name in SHEET_NAMES}
    for record in records:
        sheet_name = str(record.get("_sheet") or "")
        if sheet_name in grouped:
            grouped[sheet_name].append(record)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in SHEET_NAMES:
        columns = DROPPED_COLUMNS if sheet_name == "DROPPED" else KEPT_COLUMNS
        widths = DROPPED_WIDTHS if sheet_name == "DROPPED" else WIDTHS
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(columns)
        for record in sorted(grouped[sheet_name], key=_record_sort_key, reverse=True):
            sheet.append([record.get(column) for column in columns])
            row = sheet.max_row
            url_column = columns.index("url") + 1
            url = str(record.get("url") or "")
            if url:
                sheet.cell(row, url_column).hyperlink = url
                sheet.cell(row, url_column).style = "Hyperlink"
            if sheet_name == "PRIORITY":
                sheet.cell(row, KEPT_COLUMNS.index("band") + 1).fill = PRIORITY_FILL
            if sheet_name == "DE Required" or "GERMAN C1" in str(record.get("language_flags") or ""):
                if sheet_name != "DROPPED":
                    sheet.cell(row, KEPT_COLUMNS.index("language_flags") + 1).fill = AMBER_FILL
        _style_sheet(sheet, widths)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = xlsx_path.with_name(f".{xlsx_path.stem}.tmp.xlsx")
    workbook.save(temporary_path)
    temporary_path.replace(xlsx_path)


def merge_workbooks(source_paths: Iterable[str | Path], destination: str | Path) -> Path:
    destination = Path(destination)
    paths = [Path(path) for path in source_paths]
    if destination.exists() and destination not in paths:
        paths.insert(0, destination)
    records: list[dict[str, object]] = []
    for path in paths:
        if path.exists():
            records.extend(_read_records(path))
    _write_records_workbook(_deduplicate_records(records), destination)
    return destination


def _target_master_group(record: Mapping[str, object]) -> str:
    employment_format = str(record.get("employment_format") or "")
    if employment_format:
        return "vollzeit" if employment_format == "FULL-TIME / VOLLZEIT" else "other"
    source_group = str(record.get("_master_group") or "")
    return source_group if source_group in {"vollzeit", "other"} else "other"


def _rebuild_master_pair(
    output_dir: Path,
    new_sources: Iterable[tuple[Path, str]] = (),
    keywords: Mapping[str, Any] | None = None,
) -> tuple[Path, Path]:
    destinations = {
        "vollzeit": output_dir / "jobs_vollzeit.xlsx",
        "other": output_dir / "jobs_other.xlsx",
    }
    records: list[dict[str, object]] = []
    for group, destination in destinations.items():
        if destination.exists():
            records.extend(_read_records(destination, group))
    for path, group in new_sources:
        if path.exists():
            records.extend(_read_records(path, group))
    active_keywords = keywords or load_keywords(DEFAULT_KEYWORDS_PATH)
    unique = [_reclassify_record(record, active_keywords) for record in _deduplicate_records(records)]
    for group, destination in destinations.items():
        _write_records_workbook(
            (record for record in unique if _target_master_group(record) == group),
            destination,
        )
    return destinations["vollzeit"], destinations["other"]


def _daily_sort_key(path: Path) -> tuple[date, int]:
    match = DAILY_WORKBOOK_PATTERN.match(path.name)
    if not match:
        return date.min, 0
    return date.fromisoformat(match.group(2)), int(match.group(3) or 1)


def consolidate_existing_outputs(
    output_dir: str | Path,
    keywords: Mapping[str, Any] | None = None,
) -> tuple[Path, Path]:
    output_dir = Path(output_dir)
    run_dir = output_dir / "run"
    run_dir.mkdir(parents=True, exist_ok=True)
    daily_paths = [
        path
        for path in output_dir.glob("*.xlsx")
        if DAILY_WORKBOOK_PATTERN.match(path.name)
    ]
    sources = sorted(daily_paths, key=_daily_sort_key)
    tagged_sources = [
        (path, DAILY_WORKBOOK_PATTERN.match(path.name).group(1).lower())
        for path in sources
    ]
    destinations = _rebuild_master_pair(output_dir, tagged_sources, keywords)

    for path in daily_paths:
        target = run_dir / path.name
        if target.exists():
            raise FileExistsError(f"Cannot archive {path}; {target} already exists")
        shutil.move(str(path), str(target))
    return destinations


def _deduplicate_scored_jobs(items: Iterable[ScoredJob]) -> list[ScoredJob]:
    unique: dict[str, ScoredJob] = {}
    for item in items:
        record = {
            "url": item.ad.url,
            "title": item.ad.title,
            "employer": item.ad.employer,
            "location": item.ad.location,
        }
        key = _record_key(record)
        existing = unique.get(key)
        if existing is None or _item_sort_key(item) >= _item_sort_key(existing):
            unique[key] = item
    return list(unique.values())


def write_outputs(
    jobs: Iterable[ScoredJob],
    output_dir: str | Path,
    run_date: date | None = None,
    keywords: Mapping[str, Any] | None = None,
) -> tuple[Path, Path]:
    run_date = run_date or date.today()
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    run_dir = output_dir / "run"
    run_dir.mkdir(parents=True, exist_ok=True)
    items = _deduplicate_scored_jobs(jobs)
    full_time = [item for item in items if item.result.employment_format == "FULL-TIME / VOLLZEIT"]
    other = [item for item in items if item.result.employment_format != "FULL-TIME / VOLLZEIT"]
    base_stem = run_date.isoformat()
    stem = base_stem
    sequence = 2
    while any(
        (output_dir / name).exists()
        for name in (
            f"run/jobs_vollzeit_{stem}.xlsx",
            f"run/jobs_other_{stem}.xlsx",
        )
    ):
        stem = f"{base_stem}_{sequence}"
        sequence += 1
    full_time_path = run_dir / f"jobs_vollzeit_{stem}.xlsx"
    other_path = run_dir / f"jobs_other_{stem}.xlsx"
    _write_workbook(full_time, full_time_path, run_date)
    _write_workbook(other, other_path, run_date)
    _rebuild_master_pair(output_dir, ((full_time_path, "vollzeit"), (other_path, "other")), keywords)
    return full_time_path, other_path
